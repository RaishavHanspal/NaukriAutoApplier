#!/usr/bin/env python3
"""
Naukri assistant - one-shot daily pipeline.

  python naukri_assistant.py login     # local only: log in by hand, saves session.json
  python naukri_assistant.py run       # scrape -> score -> tailor -> apply (N per profile, time-capped) -> reports
  python naukri_assistant.py report    # rebuild applied_jobs.md / .csv from applied_log.json
  python naukri_assistant.py scrape | score | tailor   # individual steps (offline-friendly)

Environment:
  NAUKRI_SESSION_B64   base64 of session.json (GitHub secret). Used when session.json is absent.
  ANTHROPIC_API_KEY    optional; lets Claude answer free-text screening questions from your master resume.
  DRY_RUN=1            do everything except click Apply.
  CI=true              headless browser.

Selectors on Naukri change often. All of them live in SEL below.
"""
import base64
import csv
import json
import os
import random
import re
import sys
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path
from urllib.parse import quote

from docx import Document
from docx.shared import Pt

BASE = Path(__file__).parent
SESSION = BASE / "session.json"
JOBS_FILE = BASE / "jobs.json"
LOG_FILE = BASE / "applied_log.json"
UNANSWERED_FILE = BASE / "unanswered_questions.json"
REPORT_MD = BASE / "applied_jobs.md"
REPORT_CSV = BASE / "applied_jobs.csv"
OUT = BASE / "resumes"
DEBUG = BASE / "debug"

IST = timezone(timedelta(hours=5, minutes=30))
DRY = os.environ.get("DRY_RUN", "").lower() in ("1", "true", "yes")
HEADLESS = os.environ.get("CI", "").lower() in ("1", "true", "yes")


def load_json(path, default):
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else default


def save_json(path, data):
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")


CFG = load_json(BASE / "config.json", {})
MASTER = load_json(BASE / "master_resume.json", {})
ANSWERS = load_json(BASE / "answers.json", {"rules": []})


def now_ist():
    return datetime.now(IST)


START = time.monotonic()


def minutes_used():
    return (time.monotonic() - START) / 60


def time_left():
    """Minutes left of max_run_minutes (the whole run, scraping included)."""
    return CFG.get("max_run_minutes", 45) - minutes_used()


class Tee:
    """Mirror stdout into run.log so the bat run leaves a log that gets committed."""
    def __init__(self, path):
        self.f = open(path, "w", encoding="utf-8")
        self.out = sys.stdout

    def write(self, s):
        self.out.write(s)
        self.f.write(s)
        self.f.flush()

    def flush(self):
        self.out.flush()
        self.f.flush()


def pause(a=3.0, b=7.0):
    time.sleep(random.uniform(a, b))


def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


# ------------------------------------------------------------------ selectors (verify in DevTools)
SEL = {
    "card": ".srp-jobtuple-wrapper",
    "jd": '[class*="dang-inner-html"]',
    "apply_btn": '#apply-button, button:has-text("Apply"):not(:has-text("company site")):not(:has-text("Applied"))',
    "company_btn": '#company-site-button, button:has-text("company site"), a:has-text("company site")',
    "already": '#already-applied, button:text-is("Applied"), :text-matches("^\\\\s*Already Applied\\\\s*$", "i")',
    "success": "text=/successfully applied|application (has been )?(sent|submitted)|you have applied/i",
    "chat_root": '[class*="chatbot_DrawerContentWrapper"], [class*="chatbot_Drawer"]',
    "bot_msg": '[class*="botMsg"]',
    "chip": '[class*="chatbot_Chip"], [class*="ssrc__label"], [class*="chatbot_ChipContainer"] label',
    "text_input": '[class*="chatbot_InputContainer"] [contenteditable="true"], '
                  '[class*="chatbot_InputContainer"] input, [class*="chatbot"] textarea',
    "send": '[class*="sendMsg"], button:has-text("Save"), button:has-text("Submit")',
    "resume_input": 'input[type="file"]',
    "quota": "text=/daily (apply )?(quota|limit)|reached.{0,40}(quota|limit)|cannot apply.{0,40}today/i",
}

CARD_JS = """
els => els.map(e => {
  const q = s => e.querySelector(s);
  const t = q('a.title');
  return {
    title: t ? t.innerText.trim() : '',
    url: t ? t.href.split('?')[0] : '',
    company: q('.comp-name') ? q('.comp-name').innerText.trim() : '',
    experience: q('.exp-wrap') ? q('.exp-wrap').innerText.trim() : '',
    location: q('.loc-wrap') ? q('.loc-wrap').innerText.trim() : '',
    snippet: q('.job-desc') ? q('.job-desc').innerText.trim() : '',
  };
})
"""


class Blocked(Exception):
    pass


class SessionExpired(Blocked):
    pass


class QuotaReached(Exception):
    pass


# ------------------------------------------------------------------ browser / session
def ensure_session():
    b64 = os.environ.get("NAUKRI_SESSION_B64", "").strip()
    if b64 and not SESSION.exists():
        SESSION.write_bytes(base64.b64decode(b64))
    if not SESSION.exists():
        sys.exit("No session. Run `python naukri_assistant.py login` locally, "
                 "or set the NAUKRI_SESSION_B64 secret.")


def new_context(p):
    ensure_session()
    browser = p.chromium.launch(headless=HEADLESS, args=["--disable-blink-features=AutomationControlled"])
    ctx = browser.new_context(
        storage_state=str(SESSION),
        viewport={"width": 1366, "height": 850},
        locale="en-IN",
        timezone_id="Asia/Kolkata",
        user_agent=("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                    "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"),
    )
    return browser, ctx


def goto_checked(page, url):
    page.goto(url, wait_until="domcontentloaded")
    title = (page.title() or "").lower()
    if "access denied" in title or "captcha" in page.url.lower():
        screenshot(page, "blocked")
        raise Blocked(f"Naukri blocked this browser (bot detection) at {page.url} | page title: {page.title()!r}")
    if "nlogin" in page.url:
        raise SessionExpired("Session expired - run `login` again and refresh the NAUKRI_SESSION_B64 secret")


def login():
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False)
        ctx = browser.new_context()
        page = ctx.new_page()
        page.goto("https://www.naukri.com/nlogin/login")
        input("Log in in the browser (OTP/captcha included), then press Enter here... ")
        ctx.storage_state(path=str(SESSION))
        browser.close()
    print(f"Saved {SESSION}. For GitHub Actions run:  base64 -w0 session.json  and store it as NAUKRI_SESSION_B64")


# ------------------------------------------------------------------ applied log (dedupe + report data)
DONE = {"applied", "already_applied", "company_site", "unconfirmed"}


def load_log():
    return load_json(LOG_FILE, [])


def migrate_log():
    """First run: seed the log from the old `applied` flags in jobs.json."""
    if LOG_FILE.exists():
        return
    jobs = load_json(JOBS_FILE, {})
    log = []
    for url, j in jobs.items():
        if j.get("applied"):
            rp = j.get("resume_path", "")
            rp = ("resumes/" + re.split(r"[\\/]", rp)[-1]) if rp else ""
            log.append({"date": "before-automation", "url": url, "title": j["title"], "company": j["company"],
                        "score": j.get("match", {}).get("score"), "resume": rp, "status": "applied",
                        "detail": "migrated from jobs.json", "attempts": 1})
    save_json(LOG_FILE, log)
    print(f"Migrated {len(log)} previously applied jobs into {LOG_FILE.name}")


def log_result(j, status, detail=""):
    log = load_log()
    prev = [e for e in log if e["url"] == j["url"]]
    attempts = (prev[-1].get("attempts", 1) + 1) if prev else 1
    log = [e for e in log if e["url"] != j["url"]]
    log.append({"date": now_ist().strftime("%Y-%m-%d %H:%M"), "url": j["url"], "title": j["title"],
                "company": j["company"], "location": j.get("location", ""),
                "score": j.get("match", {}).get("score"),
                "resume": j.get("resume_path", ""), "status": status, "detail": detail, "attempts": attempts})
    save_json(LOG_FILE, log)


def unanswered_options(question):
    for d in load_json(UNANSWERED_FILE, []):
        if d["question"] == question:
            return d.get("options", [])
    return []


def now_answerable(e):
    """A job skipped for a question we can answer now (new answers.json rule / fallback) gets retried."""
    m = re.match(r"(no answer for|stuck on): (.+)", e.get("detail", ""), re.S)
    return bool(e["status"] == "needs_manual" and m
                and answer_offline(m.group(2), unanswered_options(m.group(2))) is not None)


def blocked_urls():
    """Jobs we must not try again."""
    out = set()
    for e in load_log():
        if e["status"] in DONE:
            out.add(e["url"])
        elif e.get("attempts", 1) >= CFG.get("max_attempts_per_job", 2) and not now_answerable(e):
            out.add(e["url"])
    return out


def prune_unanswered():
    """Drop questions from unanswered_questions.json that answers.json / the fallback now handles."""
    data = load_json(UNANSWERED_FILE, [])
    keep = [d for d in data if answer_offline(d["question"], d.get("options", [])) is None]
    if len(keep) != len(data):
        save_json(UNANSWERED_FILE, keep)
        print(f"{len(data) - len(keep)} previously unanswered questions can now be answered")


def applied_today():
    d = now_ist().strftime("%Y-%m-%d")
    return sum(1 for e in load_log() if e["status"] == "applied" and e["date"].startswith(d))


# ------------------------------------------------------------------ scrape
def search(page, jobs, kw, pages):
    """Scrape result pages `pages` (a range) for one profile keyword into `jobs`."""
    loc = CFG["location"]
    today = now_ist().strftime("%Y-%m-%d")
    for pg in pages:
        suffix = "" if pg == 1 else f"-{pg}"
        url = (f"https://www.naukri.com/{slug(kw)}-jobs-in-{slug(loc)}{suffix}"
               f"?k={quote(kw)}&l={quote(loc)}")
        print("Search:", url)
        goto_checked(page, url)
        try:
            page.wait_for_selector(SEL["card"], timeout=15000)
        except Exception:
            print("  no cards (end of results or layout change)")
            break
        for c in page.eval_on_selector_all(SEL["card"], CARD_JS):
            if not c["url"]:
                continue
            j = jobs.setdefault(c["url"], c)
            if kw not in j.setdefault("profiles", []):
                j["profiles"].append(kw)
            j["last_seen"] = today
        pause()


def fetch_descriptions(page, jobs, deadline, profile=None):
    skip = blocked_urls()
    max_exp = CFG.get("max_min_experience", 3)
    # only fetch JDs for jobs we could actually apply to; freshest listings first
    pending = [j for j in jobs.values() if not j.get("description") and j["url"] not in skip
               and min_exp(j) <= max_exp and (not profile or profile in job_profiles(j))]
    pending.sort(key=lambda j: j.get("last_seen", ""), reverse=True)
    for j in pending[: CFG.get("max_new_descriptions", 40)]:
        if minutes_used() >= deadline:
            print("Scrape time budget used up - continuing with the JDs fetched so far")
            break
        goto_checked(page, j["url"])
        try:
            page.wait_for_selector(SEL["jd"], timeout=10000)
            j["description"] = page.inner_text(SEL["jd"])
        except Exception:
            j["description"] = j.get("snippet", "")
        print("  JD:", j["title"], "-", j["company"])
        save_json(JOBS_FILE, jobs)
        pause()
    save_json(JOBS_FILE, jobs)
    print(f"{len(jobs)} jobs known")


# ------------------------------------------------------------------ score
def master_text():
    parts = [MASTER.get("summary", ""), " ".join(MASTER.get("skills", []))]
    for e in MASTER["experience"]:
        parts += [e["title"]] + e["bullets"]
    return " ".join(parts).lower()


def has_term(term, text):
    return re.search(r"(?<![a-z0-9+#])" + re.escape(term.lower()) + r"(?![a-z0-9+#])", text) is not None


def score_job(j):
    jd = (j.get("title", "") + " " + j.get("description", "")).lower()
    jd_skills = [s for s in CFG["skills_vocab"] if has_term(s, jd)]
    mtxt = master_text()
    matched = [s for s in jd_skills if has_term(s, mtxt)]
    missing = [s for s in jd_skills if s not in matched]
    pct = round(100 * len(matched) / len(jd_skills)) if len(jd_skills) >= CFG["min_jd_skills"] else 0
    area = any(a in (j.get("location", "") + jd).lower() for a in CFG["area_terms"])
    return {"score": pct, "matched": matched, "missing": missing, "area_match": area}


def min_exp(j):
    m = re.match(r"\s*(\d+)\s*-\s*(\d+)", j.get("experience", ""))
    return int(m.group(1)) if m else 0


def score_all(quiet=False):
    jobs = load_json(JOBS_FILE, {})
    for j in jobs.values():
        j["match"] = score_job(j)
    save_json(JOBS_FILE, jobs)
    if quiet:
        return
    ranked = sorted(jobs.values(), key=lambda r: r["match"]["score"], reverse=True)
    for r in ranked[:25]:
        print(f"{r['match']['score']:>3}%  {r['title'][:40]:40} {r['company'][:28]:28} {r.get('experience','')}")


def job_profiles(j):
    """Profile keywords a job belongs to (where it was found; older jobs: by title)."""
    if j.get("profiles"):
        return j["profiles"]
    return [kw for kw in CFG["keywords"] if has_term(kw, j.get("title", "").lower())]


def candidates(profile=None):
    jobs = load_json(JOBS_FILE, {})
    skip = blocked_urls()
    out = []
    for j in jobs.values():
        m = j.get("match") or {}
        if j["url"] in skip or m.get("score", 0) < CFG["threshold"]:
            continue
        if min_exp(j) > CFG.get("max_min_experience", 3):
            continue
        if CFG.get("require_area_match") and not m.get("area_match"):
            continue
        if profile and profile not in job_profiles(j):
            continue
        out.append(j)
    # listings seen in this/recent searches are most likely still open
    out.sort(key=lambda j: (j.get("last_seen", ""), j["match"]["area_match"], j["match"]["score"]), reverse=True)
    return out


# ------------------------------------------------------------------ tailor
def build_resume(j, path):
    jd = (j or {}).get("description", "").lower()
    jd_skills = set((j or {}).get("match", {}).get("matched", []))
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Pt(36)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.add_heading(MASTER["name"], 0)
    doc.add_paragraph(" | ".join(MASTER.get("contact", [])))
    doc.add_heading("Summary", 1)
    doc.add_paragraph(MASTER["summary"])
    skills = sorted(MASTER["skills"], key=lambda s: (s.lower() not in jd_skills, not has_term(s, jd)))
    doc.add_heading("Skills", 1)
    doc.add_paragraph(", ".join(skills))
    doc.add_heading("Experience", 1)
    for e in MASTER["experience"]:
        doc.add_paragraph().add_run(f"{e['title']} - {e['company']} ({e['dates']})").bold = True
        bullets = sorted(e["bullets"], key=lambda b: -sum(has_term(s, b.lower()) for s in jd_skills))
        for b in bullets:
            doc.add_paragraph(b, style="List Bullet")
    doc.add_heading("Education", 1)
    for ed in MASTER.get("education", []):
        doc.add_paragraph(ed)
    doc.save(path)


def tailor_for(j):
    OUT.mkdir(exist_ok=True)
    fn = OUT / f"{slug(MASTER['name'])}_{slug(j['company'])}_{slug(j['title'])[:40]}.docx"
    build_resume(j, fn)
    j["resume_path"] = f"resumes/{fn.name}"
    return fn


def tailor_all():
    jobs = load_json(JOBS_FILE, {})
    n = 0
    for j in jobs.values():
        if j.get("match", {}).get("score", 0) >= CFG["threshold"]:
            tailor_for(j)
            n += 1
    save_json(JOBS_FILE, jobs)
    print(f"{n} tailored resumes in {OUT}")


def upload_profile_resume(page, path):
    """Naukri applies with the resume on your profile, so swap it before applying."""
    goto_checked(page, "https://www.naukri.com/mnjuser/profile")
    pause(2, 4)
    page.locator(SEL["resume_input"]).first.set_input_files(str(path))
    pause(4, 6)


# ------------------------------------------------------------------ questionnaire answering
def usable(answer):
    if answer is None:
        return False
    items = answer if isinstance(answer, list) else [answer]
    return all(a and "REPLACE_ME" not in str(a) for a in items)


def match_option(candidates_, options):
    cands = candidates_ if isinstance(candidates_, list) else [candidates_]
    low = {o.lower().strip(): o for o in options}
    for c in cands:
        c = str(c).lower().strip()
        if c in low:
            return low[c]
    for c in cands:
        c = str(c).lower().strip()
        for k, o in low.items():
            if c and (c in k or k in c):
                return o
    return None


MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct",
                                       "nov", "dec"], 1)}
STOP = {"in", "on", "of", "with", "and", "or", "the", "a", "an", "any", "your", "tools", "tool", "domain",
        "field", "role", "work", "working", "experience", "hands", "relevant", "using", "as"}


def stems(text):
    """Crude stemming so 'analytics'/'analysis'/'analysed' and 'reporting'/'reports' match."""
    return {w[:5] if len(w) >= 5 else w for w in re.findall(r"[a-z0-9+#]+", text.lower()) if w not in STOP}


def role_months(e):
    d = re.findall(r"([A-Za-z]{3})[a-z]*\.?\s+(\d{4})", e.get("dates", ""))
    if len(d) < 2:
        return 0
    (m1, y1), (m2, y2) = d[0], d[-1]
    return (int(y2) - int(y1)) * 12 + MONTHS.get(m2.lower(), 1) - MONTHS.get(m1.lower(), 1) + 1


def skill_years(skill):
    """Years across the dated resume roles whose text mentions every word of `skill`. 0 if none."""
    want = stems(skill)
    if not want:
        return None
    months = sum(role_months(e) for e in MASTER.get("experience", [])
                 if want <= stems(" ".join([e["title"]] + e["bullets"])))
    return round(months / 12, 1)


def pick_years_option(years, options):
    """Choose the option range ('1-2 Years', '3+ years', 'Less than 1 year', 'Fresher') containing `years`."""
    for o in options:
        ol = o.lower()
        nums = [float(x) for x in re.findall(r"\d+(?:\.\d+)?", ol)]
        if not nums:
            if years < 0.5 and re.search(r"fresher|no experience|none|^no$|never", ol):
                return o
            continue
        if len(nums) >= 2:
            lo, hi = nums[0], nums[1]
        elif re.search(r"\+|more|above|over|greater|or more", ol):
            lo, hi = nums[0], float("inf")
        elif re.search(r"less|below|under|<|upto|up to", ol):
            lo, hi = 0, nums[0] - 1e-9
        else:
            lo = hi = nums[0]
        if "month" in ol and "year" not in ol:
            lo, hi = lo / 12, hi / 12
        if lo <= years <= hi or (lo == hi and int(years) == lo):
            return o
    return None


def fallback_answer(question, options):
    """Answer skill-experience questions from the master resume only. None if not that kind of question."""
    q = question.strip().rstrip("?").strip()
    m = re.search(r"(?:how many|no\.? of|number of)\s+years?.{0,40}?experience.{0,30}?\b(?:in|with|on|of|as)\s+(.+)$",
                  q, re.I)
    if m:
        years = skill_years(m.group(1))
        if years is None:
            return None
        if options:
            return pick_years_option(years, options)
        return f"{years:g}"
    m = re.search(r"^(?:do|have|did) you (?:have |had )?(?:any |prior |practical |working |hands[- ]on |relevant )*"
                  r"(?:experience|exposure|worked|knowledge)\s*(?:in|with|on|of)?\s+(.+)$", q, re.I)
    if m and (not options or match_option(["Yes", "No"], options)):
        want = stems(m.group(1))
        if not want:
            return None
        resume = stems(master_text())
        ans = "Yes" if want <= resume else "No"
        return match_option(ans, options) if options else ans
    return None


def llm_answer(question, options):
    key = os.environ.get("ANTHROPIC_API_KEY")
    if not key or not CFG.get("use_llm_for_questions", True):
        return None
    try:
        import anthropic
        client = anthropic.Anthropic(api_key=key)
        facts = [r for r in ANSWERS.get("rules", []) if usable(r.get("answer"))]
        system = (
            "You answer job-application screening questions on behalf of the candidate, using ONLY the "
            "facts in the master resume and profile facts provided. Never invent facts, numbers, salary, "
            "dates, employers or skills. You may infer years of use of a tool only from the dated roles in "
            "the resume. If the answer is not clearly supported, reply with exactly UNKNOWN. "
            "If options are given, reply with exactly one option copied verbatim. For free text reply in "
            "1-3 plain first-person sentences, no markdown.")
        user = (f"MASTER RESUME:\n{json.dumps(MASTER)}\n\nPROFILE FACTS:\n{json.dumps(facts)}\n\n"
                f"QUESTION: {question}\nOPTIONS: {options if options else 'none (free text)'}")
        msg = client.messages.create(model=CFG.get("llm_model", "claude-haiku-5-5"), max_tokens=300,
                                     system=system, messages=[{"role": "user", "content": user}])
        text = msg.content[0].text.strip()
        if text.upper().startswith("UNKNOWN"):
            return None
        return match_option(text, options) if options else text
    except Exception as ex:
        print("  LLM answer failed:", ex)
        return None


def rule_answer(question, options):
    ql = question.lower().strip().rstrip("?").strip()
    for rule in ANSWERS.get("rules", []):
        if re.search(rule["pattern"], ql) and usable(rule.get("answer")):
            a = rule["answer"]
            if options:
                m = match_option(a, options)
                if m:
                    return m
            else:
                return a[0] if isinstance(a, list) else a
    return None


def answer_offline(question, options):
    """answers.json rules, then the resume-based fallback. No network."""
    a = rule_answer(question, options)
    return a if a is not None else fallback_answer(question, options)


def answer_question(question, options):
    """Returns an answer string or None (None = we don't know, so a human must answer)."""
    a = rule_answer(question, options)
    if a is None:
        a = llm_answer(question, options)
    if a is None:
        a = fallback_answer(question, options)
    return a


def record_unanswered(j, question, options):
    data = load_json(UNANSWERED_FILE, [])
    if not any(d["question"] == question for d in data):
        data.append({"question": question, "options": options, "example_job": j["url"],
                     "how_to_fix": "add a rule for this to answers.json"})
        save_json(UNANSWERED_FILE, data)


# ------------------------------------------------------------------ apply
def confirm_applied(page, url):
    if page.locator(SEL["success"]).count():
        return True
    goto_checked(page, url)
    pause(1.5, 3)
    return page.locator(SEL["already"]).count() > 0


def run_questionnaire(page, j):
    last_q, same = None, 0
    for _ in range(CFG.get("max_questions", 15)):
        if page.locator(SEL["success"]).count():
            return "applied", ""
        root = page.locator(SEL["chat_root"]).first
        if not page.locator(SEL["chat_root"]).count():
            break
        msgs = root.locator(SEL["bot_msg"])
        n = msgs.count()
        question = msgs.nth(n - 1).inner_text().strip() if n else ""
        options = [c.inner_text().strip() for c in root.locator(SEL["chip"]).all()]
        options = [o for o in options if o]
        if not question:
            pause(1, 2)
            continue
        same = same + 1 if question == last_q else 0
        last_q = question
        if same >= 2:
            record_unanswered(j, question, options)
            return "needs_manual", f"stuck on: {question}"
        ans = answer_question(question, options)
        if ans is None:
            record_unanswered(j, question, options)
            return "needs_manual", f"no answer for: {question}"
        print(f"  Q: {question[:80]} -> {ans[:60]}")
        try:
            if options:
                for part in [a.strip() for a in ans.split("||")]:
                    root.locator(SEL["chip"]).filter(
                        has_text=re.compile(rf"^\s*{re.escape(part)}\s*$", re.I)).first.click()
                if root.locator(SEL["send"]).count():
                    root.locator(SEL["send"]).first.click()
            else:
                box = root.locator(SEL["text_input"]).first
                box.click()
                box.fill(ans)
                if root.locator(SEL["send"]).count():
                    root.locator(SEL["send"]).first.click()
                else:
                    page.keyboard.press("Enter")
        except Exception as ex:
            return "needs_manual", f"could not answer '{question}': {ex}"
        pause(1.5, 3)
    return ("applied", "") if page.locator(SEL["success"]).count() else ("unconfirmed", "")


def apply_to(page, j):
    goto_checked(page, j["url"])
    pause(2, 4)
    if page.locator(SEL["already"]).count():
        return "already_applied", ""
    has_apply = page.locator(SEL["apply_btn"]).count() > 0
    if not has_apply:
        if page.locator(SEL["company_btn"]).count():
            return "company_site", "Apply on company site - open the link and apply manually"
        return "failed", "no Apply button found (layout change or expired listing)"
    if DRY:
        return "dry_run", "would apply"
    page.locator(SEL["apply_btn"]).first.click()
    pause(2.5, 4.5)
    if page.locator(SEL["quota"]).count():
        raise QuotaReached(page.locator(SEL["quota"]).first.inner_text()[:150])
    status, detail = run_questionnaire(page, j)
    if status in ("applied", "unconfirmed"):
        status = "applied" if confirm_applied(page, j["url"]) else "unconfirmed"
    return status, detail


def screenshot(page, name):
    try:
        DEBUG.mkdir(exist_ok=True)
        page.screenshot(path=str(DEBUG / f"{slug(name)[:60]}.png"))
    except Exception:
        pass


def apply_loop(page, target):
    """Apply to up to `target` jobs per profile keyword, round-robin across profiles, until time runs out.
    Returns {profile: applied count}."""
    profiles = list(CFG["keywords"])
    queues = {kw: candidates(kw) for kw in profiles}
    applied = {kw: 0 for kw in profiles}
    tried = set()
    clicks = 0
    master_docx = OUT / "_master_resume.docx"
    swapped = False
    try:
        while True:
            active = [kw for kw in profiles if applied[kw] < target and queues[kw]]
            if not active:
                break
            for kw in active:
                if time_left() < 2:  # keep time to restore the master resume
                    print(f"\nTime limit ({CFG.get('max_run_minutes', 45)} min) almost reached after "
                          f"{minutes_used():.0f} min - stopping and restoring the master resume")
                    return applied
                if clicks >= CFG.get("max_clicks_per_run", 90):
                    print("Safety cap reached for this run (max_clicks_per_run)")
                    return applied
                j = queues[kw].pop(0)
                if j["url"] in tried:  # same job listed under several profiles
                    continue
                tried.add(j["url"])
                print(f"\n<{kw} {applied[kw]}/{target}> [{j['match']['score']}%] {j['title']} @ {j['company']} "
                      f"({j.get('experience','')}) | {minutes_used():.0f} min used")
                resume = tailor_for(j)
                try:
                    if CFG.get("upload_tailored_resume", True) and not DRY:
                        upload_profile_resume(page, resume)
                        swapped = True
                    status, detail = apply_to(page, j)
                except (Blocked, QuotaReached):
                    raise
                except Exception as ex:
                    status, detail = "failed", str(ex)[:200]
                print("  ->", status, detail)
                if status == "dry_run":
                    continue
                clicks += 1 if status in ("applied", "unconfirmed", "needs_manual") else 0
                if status in ("failed", "needs_manual", "unconfirmed"):
                    screenshot(page, f"{j['company']}-{j['title']}")
                log_result(j, status, detail)
                if status == "applied":
                    applied[kw] += 1
                pause(*CFG.get("apply_delay", [10, 25]))
    except QuotaReached as ex:
        print("\nNaukri's daily apply quota reached - stopping:", ex)
    finally:
        restore_master(page, swapped, master_docx)
    return applied


def restore_master(page, swapped, master_docx):
    if swapped and CFG.get("restore_master_resume", True):
        OUT.mkdir(exist_ok=True)
        build_resume(None, master_docx)
        try:
            upload_profile_resume(page, master_docx)
            print("Profile resume restored to master")
        except Exception as ex:
            print("WARNING: could not restore master resume on profile:", ex)


# ------------------------------------------------------------------ reports
def write_reports():
    log = load_log()
    applied = sorted([e for e in log if e["status"] == "applied"], key=lambda e: e["date"], reverse=True)

    def link(e):
        return f"[{Path(e['resume']).name}]({quote(e['resume'])})" if e.get("resume") else "-"

    lines = ["# Applied jobs", "", f"Total applied: **{len(applied)}**  |  last updated "
             f"{now_ist().strftime('%Y-%m-%d %H:%M IST')}", "",
             "| Date | Role | Company | Match | Resume used | Job |", "|---|---|---|---|---|---|"]
    for e in applied:
        lines.append(f"| {e['date']} | {e['title']} | {e['company']} | {e.get('score','')}% | "
                     f"{link(e)} | [open]({e['url']}) |")
    comp = [e for e in log if e["status"] == "company_site"]
    if comp:
        lines += ["", "## Apply manually: company-site jobs", "", "| Role | Company | Match | Resume ready | Job |",
                  "|---|---|---|---|---|"]
        lines += [f"| {e['title']} | {e['company']} | {e.get('score','')}% | {link(e)} | [open]({e['url']}) |"
                  for e in comp]
    man = [e for e in log if e["status"] in ("needs_manual", "unconfirmed", "failed")]
    if man:
        lines += ["", "## Needs your attention", "", "| Role | Company | Status | Detail | Job |", "|---|---|---|---|---|"]
        lines += [f"| {e['title']} | {e['company']} | {e['status']} | {e.get('detail','')[:100]} | "
                  f"[open]({e['url']}) |" for e in man]
    REPORT_MD.write_text("\n".join(lines) + "\n", encoding="utf-8")
    with REPORT_CSV.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["date", "title", "company", "location", "match_pct", "status", "resume", "url", "detail"])
        for e in sorted(log, key=lambda e: e["date"], reverse=True):
            w.writerow([e["date"], e["title"], e["company"], e.get("location", ""), e.get("score", ""),
                        e["status"], e.get("resume", ""), e["url"], e.get("detail", "")])
    print(f"Reports written: {REPORT_MD.name}, {REPORT_CSV.name}")


# ------------------------------------------------------------------ one-go pipeline
def scrape_profiles(page, target, budget):
    """Search every profile keyword; dig deeper for profiles with fewer than `target` eligible jobs.
    Stops scraping after `budget` minutes so there is time left to apply."""
    jobs = load_json(JOBS_FILE, {})
    first = CFG.get("pages", 2)
    for kw in CFG["keywords"]:
        if minutes_used() >= budget:
            break
        search(page, jobs, kw, range(1, first + 1))
    save_json(JOBS_FILE, jobs)
    fetch_descriptions(page, jobs, budget)
    score_all(quiet=True)
    for kw in CFG["keywords"]:
        have = len(candidates(kw))
        if have >= target or minutes_used() >= budget:
            continue
        print(f"{kw}: only {have} eligible - scraping deeper pages")
        jobs = load_json(JOBS_FILE, {})
        search(page, jobs, kw, range(first + 1, CFG.get("pages_if_short", 5) + 1))
        save_json(JOBS_FILE, jobs)
        fetch_descriptions(page, jobs, budget, profile=kw)
        score_all(quiet=True)
    print(f"\nEligible per profile (>= {CFG['threshold']}% match, not yet applied):")
    for kw in CFG["keywords"]:
        print(f"  {kw:22} {len(candidates(kw))}")


def run():
    global START
    from playwright.sync_api import sync_playwright
    START = time.monotonic()
    migrate_log()
    prune_unanswered()
    target = CFG.get("per_profile_target", 10)
    print(f"Run started {now_ist():%Y-%m-%d %H:%M IST}: up to {target} applications per profile x "
          f"{len(CFG['keywords'])} profiles, max {CFG.get('max_run_minutes', 45)} min{' (DRY RUN)' if DRY else ''}")
    with sync_playwright() as p:
        browser, ctx = new_context(p)
        page = ctx.new_page()
        try:
            scrape_profiles(page, target, min(CFG.get("scrape_budget_minutes", 15), CFG.get("max_run_minutes", 45)))
            applied = apply_loop(page, target)
            print(f"\nApplied this run: {sum(applied.values())} in {minutes_used():.0f} min "
                  f"| today total: {applied_today()}")
            for kw, n in applied.items():
                print(f"  {kw:22} {n}/{target}" + ("" if n >= target else "  (ran out of eligible jobs or time)"))
        except SessionExpired:
            raise
        except Blocked as ex:
            print("STOPPED:", ex)
            sys.exit(2)
        finally:
            write_reports()
            browser.close()


def run_with_relogin():
    try:
        run()
    except SessionExpired as ex:
        if HEADLESS or not sys.stdin.isatty():
            print("STOPPED:", ex)
            sys.exit(3)
        print("\nNaukri session expired. A browser window will open - log in there.")
        login()
        run()


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else ""
    if cmd == "login":
        login()
    elif cmd == "run":
        sys.stdout = Tee(BASE / "run.log")
        run_with_relogin()
    elif cmd == "report":
        migrate_log()
        write_reports()
    elif cmd == "score":
        score_all()
    elif cmd == "tailor":
        tailor_all()
    elif cmd == "scrape":
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            b, c = new_context(p)
            scrape_profiles(c.new_page(), CFG.get("per_profile_target", 10), CFG.get("scrape_budget_minutes", 15))
            b.close()
    else:
        print(__doc__)
